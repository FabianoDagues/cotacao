import csv
import io
import os
import unicodedata
import hashlib
import secrets
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_CEILING
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import FastAPI, Request, Depends, Form, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import select, and_, inspect, text, delete
from sqlalchemy.orm import Session, selectinload
from openpyxl import load_workbook, Workbook

from .db import Base, engine, get_db, SessionLocal
from .models import User, PasswordResetToken, Quote, QuoteSupplier, QuoteItem, QuoteResponse, PurchaseOrder, PurchaseOrderItem
from .security import hash_password, verify_password

BASE_DIR = Path(__file__).resolve().parent
app = FastAPI(title="Sistema de Cotação")
app.add_middleware(SessionMiddleware, secret_key=os.getenv("SECRET_KEY", "troque-esta-chave-em-producao"), https_only=os.getenv("COOKIE_SECURE", "0") == "1")
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")
APP_TIMEZONE = os.getenv("APP_TIMEZONE", "America/Sao_Paulo")

def now_local() -> datetime:
    return datetime.now(ZoneInfo(APP_TIMEZONE)).replace(tzinfo=None)


def parse_decimal(value, default="0") -> Decimal:
    if value is None or value == "":
        return Decimal(default)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    return Decimal(text)


def current_user(request: Request, db: Session) -> Optional[User]:
    uid = request.session.get("user_id")
    if not uid:
        return None
    return db.get(User, uid)


def require_role(request: Request, db: Session, role: str) -> User:
    user = current_user(request, db)
    if not user or not user.active or user.role != role:
        raise HTTPException(status_code=403, detail="Acesso negado")
    return user


def ensure_runtime_schema():
    """Aplica pequenas migrações compatíveis com bancos já existentes."""
    inspector = inspect(engine)
    tables = set(inspector.get_table_names())
    if "users" in tables:
        columns = {column["name"] for column in inspector.get_columns("users")}
        if "must_change_password" not in columns:
            with engine.begin() as conn:
                default_value = "FALSE" if engine.dialect.name == "postgresql" else "0"
                conn.execute(text(f"ALTER TABLE users ADD COLUMN must_change_password BOOLEAN NOT NULL DEFAULT {default_value}"))

    # Marca quando cada fornecedor efetivamente clicou em salvar/enviar a cotação.
    if "quote_suppliers" in tables:
        qs_columns = {column["name"] for column in inspector.get_columns("quote_suppliers")}
        if "submitted_at" not in qs_columns:
            with engine.begin() as conn:
                conn.execute(text("ALTER TABLE quote_suppliers ADD COLUMN submitted_at TIMESTAMP NULL"))


def backfill_submission_status():
    """Preserva o status de cotações respondidas antes da atualização que criou submitted_at."""
    with SessionLocal() as db:
        links = db.scalars(select(QuoteSupplier).where(QuoteSupplier.submitted_at.is_(None))).all()
        changed = False
        for link in links:
            item_ids = list(db.scalars(select(QuoteItem.id).where(QuoteItem.quote_id == link.quote_id)).all())
            if not item_ids:
                continue
            responses = db.scalars(
                select(QuoteResponse).where(
                    QuoteResponse.supplier_id == link.supplier_id,
                    QuoteResponse.quote_item_id.in_(item_ids),
                )
            ).all()
            if responses:
                link.submitted_at = max((r.updated_at for r in responses if r.updated_at), default=now_local())
                changed = True
        if changed:
            db.commit()


def seed_admin():
    email = os.getenv("ADMIN_EMAIL", "admin@empresa.local").lower()
    password = os.getenv("ADMIN_PASSWORD", "Admin123!")
    name = os.getenv("ADMIN_NAME", "Administrador")
    with SessionLocal() as db:
        existing = db.scalar(select(User).where(User.email == email))
        if not existing:
            db.add(User(name=name, email=email, password_hash=hash_password(password), role="admin", company_name="Administração", must_change_password=False))
            db.commit()


@app.on_event("startup")
def startup():
    Base.metadata.create_all(bind=engine)
    ensure_runtime_schema()
    # create_all é chamado novamente para garantir tabelas novas após pequenas migrações.
    Base.metadata.create_all(bind=engine)
    backfill_submission_status()
    seed_admin()


@app.get("/healthz")
def healthz():
    return {"status": "ok"}


@app.get("/", response_class=HTMLResponse)
def home(request: Request, db: Session = Depends(get_db)):
    user = current_user(request, db)
    if not user:
        return RedirectResponse("/login", 303)
    if user.role == "supplier" and user.must_change_password:
        return RedirectResponse("/fornecedor/alterar-senha", 303)
    return RedirectResponse("/admin" if user.role == "admin" else "/fornecedor", 303)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html", {"error": None})


@app.post("/login", response_class=HTMLResponse)
def login(request: Request, email: str = Form(...), password: str = Form(...), db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.email == email.strip().lower()))
    if not user or not user.active or not verify_password(password, user.password_hash):
        return templates.TemplateResponse(request, "login.html", {"error": "E-mail ou senha inválidos."}, status_code=400)
    request.session.clear()
    request.session["user_id"] = user.id
    request.session["role"] = user.role
    if user.role == "supplier" and user.must_change_password:
        return RedirectResponse("/fornecedor/alterar-senha", 303)
    return RedirectResponse("/admin" if user.role == "admin" else "/fornecedor", 303)


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", 303)


@app.get("/admin", response_class=HTMLResponse)
def admin_dashboard(request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    quotes = db.scalars(select(Quote).order_by(Quote.created_at.desc())).all()
    submission_stats = {}
    for quote in quotes:
        links = db.scalars(select(QuoteSupplier).where(QuoteSupplier.quote_id == quote.id)).all()
        total = len(links)
        sent = sum(1 for link in links if link.submitted_at is not None)
        submission_stats[quote.id] = {"total": total, "sent": sent, "pending": total - sent}
    return templates.TemplateResponse(
        request,
        "admin_dashboard.html",
        {"user": user, "quotes": quotes, "now": now_local(), "submission_stats": submission_stats},
    )


@app.get("/admin/fornecedores", response_class=HTMLResponse)
def admin_suppliers(request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    suppliers = db.scalars(select(User).where(User.role == "supplier").order_by(User.company_name, User.name)).all()
    created = request.query_params.get("criado") == "1"
    updated = request.query_params.get("editado") == "1"
    deleted = request.query_params.get("excluido") == "1"
    return templates.TemplateResponse(
        request,
        "admin_suppliers.html",
        {
            "user": user,
            "suppliers": suppliers,
            "error": None,
            "reset_link": None,
            "reset_supplier": None,
            "created": created,
            "updated": updated,
            "deleted": deleted,
        },
    )


@app.post("/admin/fornecedores")
def create_supplier(
    request: Request,
    name: str = Form(...),
    company_name: str = Form(...),
    email: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    require_role(request, db, "admin")
    normalized = email.strip().lower()
    if db.scalar(select(User).where(User.email == normalized)):
        suppliers = db.scalars(select(User).where(User.role == "supplier")).all()
        return templates.TemplateResponse(request, "admin_suppliers.html", {"user": current_user(request, db), "suppliers": suppliers, "error": "Este e-mail já está cadastrado.", "reset_link": None, "reset_supplier": None, "created": False}, status_code=400)
    db.add(User(name=name.strip(), company_name=company_name.strip(), email=normalized, password_hash=hash_password(password), role="supplier", must_change_password=True))
    db.commit()
    return RedirectResponse("/admin/fornecedores?criado=1", 303)


@app.get("/admin/fornecedores/{supplier_id}/editar", response_class=HTMLResponse)
def edit_supplier_page(supplier_id: int, request: Request, db: Session = Depends(get_db)):
    admin = require_role(request, db, "admin")
    supplier = db.get(User, supplier_id)
    if not supplier or supplier.role != "supplier":
        raise HTTPException(404, "Fornecedor não encontrado")
    return templates.TemplateResponse(
        request,
        "admin_supplier_edit.html",
        {"user": admin, "supplier": supplier, "error": None},
    )


@app.post("/admin/fornecedores/{supplier_id}/editar", response_class=HTMLResponse)
def edit_supplier_submit(
    supplier_id: int,
    request: Request,
    name: str = Form(...),
    company_name: str = Form(...),
    email: str = Form(...),
    active: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    admin = require_role(request, db, "admin")
    supplier = db.get(User, supplier_id)
    if not supplier or supplier.role != "supplier":
        raise HTTPException(404, "Fornecedor não encontrado")

    normalized = email.strip().lower()
    duplicate = db.scalar(select(User).where(User.email == normalized, User.id != supplier.id))
    if duplicate:
        return templates.TemplateResponse(
            request,
            "admin_supplier_edit.html",
            {"user": admin, "supplier": supplier, "error": "Este e-mail já está sendo usado por outro usuário."},
            status_code=400,
        )

    supplier.company_name = company_name.strip()
    supplier.name = name.strip()
    supplier.email = normalized
    supplier.active = active == "1"
    db.commit()
    return RedirectResponse("/admin/fornecedores?editado=1", 303)


@app.post("/admin/fornecedores/{supplier_id}/excluir")
def delete_supplier(supplier_id: int, request: Request, db: Session = Depends(get_db)):
    require_role(request, db, "admin")
    supplier = db.get(User, supplier_id)
    if not supplier or supplier.role != "supplier":
        raise HTTPException(404, "Fornecedor não encontrado")

    # Exclusão definitiva: remove também respostas e vínculos do fornecedor com cotações.
    # As cotações e seus produtos permanecem; apenas a participação deste fornecedor é removida.
    order_ids = list(db.scalars(select(PurchaseOrder.id).where(PurchaseOrder.supplier_id == supplier.id)).all())
    if order_ids:
        db.execute(delete(PurchaseOrderItem).where(PurchaseOrderItem.order_id.in_(order_ids)))
    db.execute(delete(PurchaseOrder).where(PurchaseOrder.supplier_id == supplier.id))
    db.execute(delete(QuoteResponse).where(QuoteResponse.supplier_id == supplier.id))
    db.execute(delete(QuoteSupplier).where(QuoteSupplier.supplier_id == supplier.id))
    db.execute(delete(PasswordResetToken).where(PasswordResetToken.user_id == supplier.id))
    db.delete(supplier)
    db.commit()
    return RedirectResponse("/admin/fornecedores?excluido=1", 303)


def token_digest(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


@app.post("/admin/fornecedores/{supplier_id}/reset-senha", response_class=HTMLResponse)
def admin_generate_password_reset(supplier_id: int, request: Request, db: Session = Depends(get_db)):
    admin = require_role(request, db, "admin")
    supplier = db.get(User, supplier_id)
    if not supplier or supplier.role != "supplier":
        raise HTTPException(404, "Fornecedor não encontrado")

    now = datetime.utcnow()
    # Invalida links anteriores ainda não utilizados para que apenas o mais recente funcione.
    previous = db.scalars(
        select(PasswordResetToken).where(
            PasswordResetToken.user_id == supplier.id,
            PasswordResetToken.used_at.is_(None),
        )
    ).all()
    for item in previous:
        item.used_at = now

    raw_token = secrets.token_urlsafe(32)
    reset = PasswordResetToken(
        user_id=supplier.id,
        token_hash=token_digest(raw_token),
        expires_at=now + timedelta(hours=1),
        created_by_user_id=admin.id,
    )
    db.add(reset)
    db.commit()

    public_base_url = os.getenv("PUBLIC_BASE_URL", "").strip().rstrip("/")
    reset_link = f"{public_base_url}/redefinir-senha/{raw_token}" if public_base_url else str(request.url_for("reset_password_page", token=raw_token))
    suppliers = db.scalars(select(User).where(User.role == "supplier").order_by(User.company_name, User.name)).all()
    return templates.TemplateResponse(
        request,
        "admin_suppliers.html",
        {
            "user": admin,
            "suppliers": suppliers,
            "error": None,
            "reset_link": reset_link,
            "reset_supplier": supplier,
            "created": False,
        },
    )


def find_valid_reset_token(db: Session, raw_token: str) -> Optional[PasswordResetToken]:
    digest = token_digest(raw_token)
    reset = db.scalar(select(PasswordResetToken).where(PasswordResetToken.token_hash == digest))
    if not reset or reset.used_at is not None or reset.expires_at < datetime.utcnow():
        return None
    return reset


@app.get("/redefinir-senha/{token}", response_class=HTMLResponse, name="reset_password_page")
def reset_password_page(token: str, request: Request, db: Session = Depends(get_db)):
    reset = find_valid_reset_token(db, token)
    supplier = db.get(User, reset.user_id) if reset else None
    return templates.TemplateResponse(
        request,
        "reset_password.html",
        {"user": None, "token": token, "valid": bool(reset and supplier and supplier.active), "supplier": supplier, "error": None, "success": False},
        status_code=200 if reset and supplier and supplier.active else 400,
    )


@app.post("/redefinir-senha/{token}", response_class=HTMLResponse)
def reset_password_submit(
    token: str,
    request: Request,
    password: str = Form(...),
    password_confirm: str = Form(...),
    db: Session = Depends(get_db),
):
    reset = find_valid_reset_token(db, token)
    supplier = db.get(User, reset.user_id) if reset else None
    if not reset or not supplier or not supplier.active:
        return templates.TemplateResponse(
            request,
            "reset_password.html",
            {"user": None, "token": token, "valid": False, "supplier": supplier, "error": "Este link expirou ou já foi utilizado.", "success": False},
            status_code=400,
        )
    if password != password_confirm:
        return templates.TemplateResponse(
            request,
            "reset_password.html",
            {"user": None, "token": token, "valid": True, "supplier": supplier, "error": "As duas senhas não são iguais.", "success": False},
            status_code=400,
        )
    if len(password) < 8:
        return templates.TemplateResponse(
            request,
            "reset_password.html",
            {"user": None, "token": token, "valid": True, "supplier": supplier, "error": "A nova senha deve ter pelo menos 8 caracteres.", "success": False},
            status_code=400,
        )

    supplier.password_hash = hash_password(password)
    supplier.must_change_password = False
    reset.used_at = datetime.utcnow()
    # Invalida qualquer outro link de redefinição que tenha sido criado para o mesmo fornecedor.
    remaining = db.scalars(
        select(PasswordResetToken).where(
            PasswordResetToken.user_id == supplier.id,
            PasswordResetToken.used_at.is_(None),
        )
    ).all()
    for item in remaining:
        item.used_at = reset.used_at
    db.commit()
    return templates.TemplateResponse(
        request,
        "reset_password.html",
        {"user": None, "token": token, "valid": False, "supplier": supplier, "error": None, "success": True},
    )


@app.get("/admin/cotacoes/nova", response_class=HTMLResponse)
def new_quote_page(request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    suppliers = db.scalars(select(User).where(and_(User.role == "supplier", User.active == True)).order_by(User.company_name)).all()
    return templates.TemplateResponse(request, "admin_quote_new.html", {"user": user, "suppliers": suppliers, "error": None})


def read_product_rows(upload: UploadFile):
    filename = (upload.filename or "").lower()
    raw = upload.file.read()
    rows = []
    if filename.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            return []
        headers = [str(x or "").strip().lower() for x in data[0]]
        for values in data[1:]:
            row = dict(zip(headers, values))
            rows.append(row)
    elif filename.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        sample = text[:2000]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        for r in reader:
            rows.append({str(k or "").strip().lower(): v for k, v in r.items()})
    else:
        raise ValueError("Envie um arquivo .xlsx ou .csv")
    return rows


def pick(row, *keys):
    for key in keys:
        if key in row and row[key] not in (None, ""):
            return row[key]
    return None


@app.post("/admin/cotacoes/nova")
def create_quote(
    request: Request,
    title: str = Form(...),
    deadline: str = Form(...),
    suppliers: list[int] = Form(default=[]),
    product_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    user = require_role(request, db, "admin")
    all_suppliers = db.scalars(select(User).where(and_(User.role == "supplier", User.active == True)).order_by(User.company_name)).all()
    try:
        deadline_dt = datetime.fromisoformat(deadline)
        rows = read_product_rows(product_file)
        if not rows:
            raise ValueError("A planilha não contém produtos.")
        if not suppliers:
            raise ValueError("Selecione ao menos um fornecedor.")
        quote = Quote(title=title.strip(), deadline=deadline_dt)
        db.add(quote)
        db.flush()
        for sid in suppliers:
            supplier = db.get(User, sid)
            if supplier and supplier.role == "supplier":
                db.add(QuoteSupplier(quote_id=quote.id, supplier_id=sid))
        imported = 0
        for row in rows:
            code = pick(row, "codigo", "código", "cod", "sku") or ""
            description = pick(row, "produto", "descricao", "descrição", "item")
            qty = pick(row, "quantidade", "qtd", "qtde")
            unit = pick(row, "unidade", "un", "und")
            if not description:
                continue
            db.add(QuoteItem(quote_id=quote.id, code=str(code).strip(), description=str(description).strip(), requested_qty=parse_decimal(qty, "1"), base_unit=str(unit or "UN").strip().upper()))
            imported += 1
        if imported == 0:
            raise ValueError("Nenhum produto válido encontrado. Use colunas: Código, Produto, Quantidade e Unidade.")
        db.commit()
        return RedirectResponse(f"/admin/cotacoes/{quote.id}", 303)
    except (ValueError, InvalidOperation) as exc:
        db.rollback()
        return templates.TemplateResponse(request, "admin_quote_new.html", {"user": user, "suppliers": all_suppliers, "error": str(exc)}, status_code=400)


@app.get("/admin/cotacoes/{quote_id}", response_class=HTMLResponse)
def admin_quote_detail(quote_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    quote = db.scalar(
        select(Quote).where(Quote.id == quote_id).options(
            selectinload(Quote.items).selectinload(QuoteItem.responses).selectinload(QuoteResponse.supplier)
        )
    )
    if not quote:
        raise HTTPException(404)

    quote_supplier_links = db.scalars(
        select(QuoteSupplier).where(QuoteSupplier.quote_id == quote_id).order_by(QuoteSupplier.id)
    ).all()
    supplier_ids = [link.supplier_id for link in quote_supplier_links]
    supplier_map = {
        supplier.id: supplier
        for supplier in (
            db.scalars(select(User).where(User.id.in_(supplier_ids))).all() if supplier_ids else []
        )
    }
    participants = []
    for link in quote_supplier_links:
        supplier = supplier_map.get(link.supplier_id)
        if supplier:
            participants.append({"supplier": supplier, "submitted_at": link.submitted_at})
    participants.sort(key=lambda row: ((row["supplier"].company_name or row["supplier"].name or "").lower()))
    suppliers = [row["supplier"] for row in participants]
    sent_participants = [row for row in participants if row["submitted_at"] is not None]
    pending_participants = [row for row in participants if row["submitted_at"] is None]

    all_active_suppliers = db.scalars(
        select(User).where(and_(User.role == "supplier", User.active == True)).order_by(User.company_name, User.name)
    ).all()
    current_ids = set(supplier_ids)
    available_suppliers = [supplier for supplier in all_active_suppliers if supplier.id not in current_ids]

    summary = []
    for item in quote.items:
        valid = [r for r in item.responses if r.normalized_price is not None]
        valid.sort(key=lambda r: r.normalized_price)
        summary.append({"item": item, "responses": valid, "min": valid[0] if valid else None, "max": valid[-1] if valid else None})

    corrected_raw = request.query_params.get("corrigido")
    try:
        corrected_response_id = int(corrected_raw) if corrected_raw else None
    except ValueError:
        corrected_response_id = None

    return templates.TemplateResponse(
        request,
        "admin_quote_detail.html",
        {
            "user": user,
            "quote": quote,
            "suppliers": suppliers,
            "participants": participants,
            "sent_participants": sent_participants,
            "pending_participants": pending_participants,
            "available_suppliers": available_suppliers,
            "summary": summary,
            "now": now_local(),
            "supplier_added": request.query_params.get("fornecedor_adicionado") == "1",
            "supplier_already_added": request.query_params.get("fornecedor_ja_adicionado") == "1",
            "corrected_response_id": corrected_response_id,
        },
    )


@app.post("/admin/cotacoes/{quote_id}/fornecedores/adicionar")
def add_supplier_to_quote(
    quote_id: int,
    request: Request,
    supplier_id: int = Form(...),
    db: Session = Depends(get_db),
):
    require_role(request, db, "admin")
    quote = db.get(Quote, quote_id)
    if not quote:
        raise HTTPException(404)

    supplier = db.get(User, supplier_id)
    if not supplier or supplier.role != "supplier" or not supplier.active:
        raise HTTPException(400, "Fornecedor inválido ou inativo")

    existing = db.scalar(
        select(QuoteSupplier).where(
            and_(QuoteSupplier.quote_id == quote_id, QuoteSupplier.supplier_id == supplier_id)
        )
    )
    if existing:
        return RedirectResponse(f"/admin/cotacoes/{quote_id}?fornecedor_ja_adicionado=1", 303)

    db.add(QuoteSupplier(quote_id=quote_id, supplier_id=supplier_id))
    db.commit()
    return RedirectResponse(f"/admin/cotacoes/{quote_id}?fornecedor_adicionado=1", 303)


@app.post("/admin/respostas/{response_id}/editar")
def admin_edit_response(
    response_id: int,
    request: Request,
    offered_unit: str = Form(...),
    factor: str = Form(...),
    offered_price: str = Form(...),
    db: Session = Depends(get_db),
):
    require_role(request, db, "admin")
    response = db.get(QuoteResponse, response_id)
    if not response:
        raise HTTPException(404)
    try:
        f = parse_decimal(factor)
        p = parse_decimal(offered_price)
        if f <= 0 or p < 0:
            raise ValueError
        response.offered_unit = offered_unit.strip().upper()
        response.base_units_per_offered_unit = f
        response.offered_price = p
        quote_id = response.item.quote_id
        db.commit()
        return RedirectResponse(f"/admin/cotacoes/{quote_id}?corrigido={response_id}#resp-{response_id}", 303)
    except Exception:
        raise HTTPException(400, "Preço ou fator inválido")


@app.post("/admin/cotacoes/{quote_id}/prazo")
def update_deadline(quote_id: int, request: Request, deadline: str = Form(...), db: Session = Depends(get_db)):
    require_role(request, db, "admin")
    quote = db.get(Quote, quote_id)
    if not quote:
        raise HTTPException(404)
    quote.deadline = datetime.fromisoformat(deadline)
    db.commit()
    return RedirectResponse(f"/admin/cotacoes/{quote_id}", 303)


@app.get("/admin/cotacoes/{quote_id}/excel")
def export_quote_excel(quote_id: int, request: Request, db: Session = Depends(get_db)):
    require_role(request, db, "admin")
    quote = db.scalar(select(Quote).where(Quote.id == quote_id).options(selectinload(Quote.items).selectinload(QuoteItem.responses).selectinload(QuoteResponse.supplier)))
    if not quote:
        raise HTTPException(404)
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparativo"
    ws.append(["Código", "Produto", "Qtd", "Unidade base", "Fornecedor", "Unidade cotada", "Fator", "Preço cotado", "Preço normalizado", "Marca", "Prazo dias"])
    for item in quote.items:
        if not item.responses:
            ws.append([item.code, item.description, float(item.requested_qty), item.base_unit])
        for r in sorted(item.responses, key=lambda x: x.normalized_price or Decimal("999999999")):
            ws.append([item.code, item.description, float(item.requested_qty), item.base_unit, r.supplier.company_name or r.supplier.name, r.offered_unit, float(r.base_units_per_offered_unit), float(r.offered_price), float(r.normalized_price) if r.normalized_price is not None else None, r.brand, r.delivery_days])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    ascii_title = unicodedata.normalize("NFKD", quote.title).encode("ascii", "ignore").decode("ascii")
    safe = "".join(c if c.isalnum() else "_" for c in ascii_title)[:50] or "cotacao"
    headers = {"Content-Disposition": f'attachment; filename="comparativo_{safe}.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)



def build_supplier_order_rows(db: Session, quote: Quote, supplier_id: int, order: Optional[PurchaseOrder]):
    """Monta as linhas do fechamento. Por padrão seleciona o menor preço de cada item."""
    existing_by_item = {item.quote_item_id: item for item in (order.items if order else [])}
    rows = []
    total = Decimal("0")
    for item in quote.items:
        supplier_response = next((r for r in item.responses if r.supplier_id == supplier_id and r.normalized_price is not None), None)
        if not supplier_response:
            continue
        valid = [r for r in item.responses if r.normalized_price is not None]
        min_response = min(valid, key=lambda r: r.normalized_price) if valid else None
        existing = existing_by_item.get(item.id)
        selected = existing is not None or (order is None and min_response is not None and min_response.id == supplier_response.id)
        factor = Decimal(supplier_response.base_units_per_offered_unit or 1)
        default_qty = (Decimal(item.requested_qty) / factor).quantize(Decimal("1"), rounding=ROUND_CEILING) if factor > 0 else Decimal(item.requested_qty)
        quantity = Decimal(existing.quantity_offered) if existing else default_qty
        line_total = (quantity * Decimal(supplier_response.offered_price)) if selected else Decimal("0")
        base_equivalent = quantity * factor
        if selected:
            total += line_total
        rows.append(
            {
                "item": item,
                "response": supplier_response,
                "is_lowest": bool(min_response and min_response.id == supplier_response.id),
                "selected": selected,
                "quantity": quantity,
                "base_equivalent": base_equivalent,
                "line_total": line_total,
            }
        )
    return rows, total


@app.get("/admin/cotacoes/{quote_id}/pedidos/{supplier_id}", response_class=HTMLResponse)
def admin_supplier_order_page(quote_id: int, supplier_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    link = db.scalar(select(QuoteSupplier).where(QuoteSupplier.quote_id == quote_id, QuoteSupplier.supplier_id == supplier_id))
    if not link:
        raise HTTPException(404, "Fornecedor não participa desta cotação")
    supplier = db.get(User, supplier_id)
    quote = db.scalar(
        select(Quote).where(Quote.id == quote_id).options(
            selectinload(Quote.items).selectinload(QuoteItem.responses).selectinload(QuoteResponse.supplier)
        )
    )
    if not quote or not supplier:
        raise HTTPException(404)
    order = db.scalar(
        select(PurchaseOrder).where(PurchaseOrder.quote_id == quote_id, PurchaseOrder.supplier_id == supplier_id).options(
            selectinload(PurchaseOrder.items)
        )
    )
    rows, total = build_supplier_order_rows(db, quote, supplier_id, order)
    return templates.TemplateResponse(
        request,
        "admin_supplier_order.html",
        {
            "user": user,
            "quote": quote,
            "supplier": supplier,
            "submission": link,
            "order": order,
            "rows": rows,
            "total": total,
            "saved": request.query_params.get("salvo") == "1",
            "finalized": request.query_params.get("finalizado") == "1",
            "error": None,
        },
    )


@app.post("/admin/cotacoes/{quote_id}/pedidos/{supplier_id}", response_class=HTMLResponse)
async def admin_supplier_order_save(quote_id: int, supplier_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    link = db.scalar(select(QuoteSupplier).where(QuoteSupplier.quote_id == quote_id, QuoteSupplier.supplier_id == supplier_id))
    supplier = db.get(User, supplier_id)
    quote = db.scalar(
        select(Quote).where(Quote.id == quote_id).options(
            selectinload(Quote.items).selectinload(QuoteItem.responses).selectinload(QuoteResponse.supplier)
        )
    )
    if not link or not supplier or not quote:
        raise HTTPException(404)
    order = db.scalar(
        select(PurchaseOrder).where(PurchaseOrder.quote_id == quote_id, PurchaseOrder.supplier_id == supplier_id).options(
            selectinload(PurchaseOrder.items)
        )
    )
    form = await request.form()
    action = str(form.get("action") or "save")
    try:
        selected_rows = []
        for item in quote.items:
            if str(form.get(f"include_{item.id}") or "") != "1":
                continue
            response = next((r for r in item.responses if r.supplier_id == supplier_id), None)
            if not response:
                continue
            qty = parse_decimal(form.get(f"qty_{item.id}"), "0")
            if qty <= 0:
                raise ValueError(f"Informe uma quantidade maior que zero para {item.description}.")
            selected_rows.append((item, response, qty))
        if action == "finalize" and not selected_rows:
            raise ValueError("Selecione ao menos um item para finalizar o pedido.")

        if not order:
            order = PurchaseOrder(quote_id=quote_id, supplier_id=supplier_id, status="draft")
            db.add(order)
            db.flush()
        else:
            db.execute(delete(PurchaseOrderItem).where(PurchaseOrderItem.order_id == order.id))

        for item, response, qty in selected_rows:
            db.add(
                PurchaseOrderItem(
                    order_id=order.id,
                    quote_item_id=item.id,
                    response_id=response.id,
                    quantity_offered=qty,
                )
            )
        order.status = "ready" if action == "finalize" else "draft"
        order.updated_at = datetime.utcnow()
        db.commit()
        flag = "finalizado=1" if action == "finalize" else "salvo=1"
        return RedirectResponse(f"/admin/cotacoes/{quote_id}/pedidos/{supplier_id}?{flag}", 303)
    except (ValueError, InvalidOperation) as exc:
        db.rollback()
        order = db.scalar(
            select(PurchaseOrder).where(PurchaseOrder.quote_id == quote_id, PurchaseOrder.supplier_id == supplier_id).options(
                selectinload(PurchaseOrder.items)
            )
        )
        rows, total = build_supplier_order_rows(db, quote, supplier_id, order)
        return templates.TemplateResponse(
            request,
            "admin_supplier_order.html",
            {
                "user": user,
                "quote": quote,
                "supplier": supplier,
                "submission": link,
                "order": order,
                "rows": rows,
                "total": total,
                "saved": False,
                "finalized": False,
                "error": str(exc),
            },
            status_code=400,
        )


@app.get("/admin/cotacoes/{quote_id}/pedidos/{supplier_id}/imprimir", response_class=HTMLResponse)
def admin_supplier_order_print(quote_id: int, supplier_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "admin")
    supplier = db.get(User, supplier_id)
    quote = db.scalar(
        select(Quote).where(Quote.id == quote_id).options(
            selectinload(Quote.items).selectinload(QuoteItem.responses).selectinload(QuoteResponse.supplier)
        )
    )
    order = db.scalar(
        select(PurchaseOrder).where(PurchaseOrder.quote_id == quote_id, PurchaseOrder.supplier_id == supplier_id).options(
            selectinload(PurchaseOrder.items)
        )
    )
    if not supplier or not quote or not order:
        raise HTTPException(404, "Salve o fechamento antes de gerar o pedido.")
    rows, total = build_supplier_order_rows(db, quote, supplier_id, order)
    selected_rows = [row for row in rows if row["selected"]]
    return templates.TemplateResponse(
        request,
        "admin_supplier_order_print.html",
        {
            "user": user,
            "quote": quote,
            "supplier": supplier,
            "order": order,
            "rows": selected_rows,
            "total": total,
            "generated_at": now_local(),
        },
    )


@app.get("/fornecedor/alterar-senha", response_class=HTMLResponse)
def supplier_change_password_page(request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "supplier")
    return templates.TemplateResponse(
        request,
        "change_password.html",
        {"user": user, "first_access": bool(user.must_change_password), "error": None},
    )


@app.post("/fornecedor/alterar-senha", response_class=HTMLResponse)
def supplier_change_password_submit(
    request: Request,
    current_password: str = Form(default=""),
    password: str = Form(...),
    password_confirm: str = Form(...),
    db: Session = Depends(get_db),
):
    user = require_role(request, db, "supplier")
    first_access = bool(user.must_change_password)

    if not first_access and not verify_password(current_password, user.password_hash):
        return templates.TemplateResponse(
            request,
            "change_password.html",
            {"user": user, "first_access": False, "error": "A senha atual está incorreta."},
            status_code=400,
        )
    if password != password_confirm:
        return templates.TemplateResponse(
            request,
            "change_password.html",
            {"user": user, "first_access": first_access, "error": "As duas senhas novas não são iguais."},
            status_code=400,
        )
    if len(password) < 8:
        return templates.TemplateResponse(
            request,
            "change_password.html",
            {"user": user, "first_access": first_access, "error": "A nova senha deve ter pelo menos 8 caracteres."},
            status_code=400,
        )
    if verify_password(password, user.password_hash):
        return templates.TemplateResponse(
            request,
            "change_password.html",
            {"user": user, "first_access": first_access, "error": "Escolha uma senha diferente da senha atual/provisória."},
            status_code=400,
        )

    user.password_hash = hash_password(password)
    user.must_change_password = False
    db.commit()
    return RedirectResponse("/fornecedor?senha=alterada", 303)


@app.get("/fornecedor", response_class=HTMLResponse)
def supplier_dashboard(request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "supplier")
    if user.must_change_password:
        return RedirectResponse("/fornecedor/alterar-senha", 303)
    quote_ids = db.scalars(select(QuoteSupplier.quote_id).where(QuoteSupplier.supplier_id == user.id)).all()
    quotes = db.scalars(select(Quote).where(Quote.id.in_(quote_ids)).order_by(Quote.deadline.desc())).all() if quote_ids else []
    password_changed = request.query_params.get("senha") == "alterada"
    return templates.TemplateResponse(request, "supplier_dashboard.html", {"user": user, "quotes": quotes, "now": now_local(), "password_changed": password_changed})


@app.get("/fornecedor/cotacoes/{quote_id}", response_class=HTMLResponse)
def supplier_quote_page(quote_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "supplier")
    if user.must_change_password:
        return RedirectResponse("/fornecedor/alterar-senha", 303)
    allowed = db.scalar(select(QuoteSupplier).where(and_(QuoteSupplier.quote_id == quote_id, QuoteSupplier.supplier_id == user.id)))
    if not allowed:
        raise HTTPException(403)
    quote = db.scalar(select(Quote).where(Quote.id == quote_id).options(selectinload(Quote.items)))
    if not quote:
        raise HTTPException(404)
    responses = db.scalars(select(QuoteResponse).where(QuoteResponse.supplier_id == user.id, QuoteResponse.quote_item_id.in_([i.id for i in quote.items]))).all() if quote.items else []
    by_item = {r.quote_item_id: r for r in responses}
    return templates.TemplateResponse(
        request,
        "supplier_quote.html",
        {
            "user": user,
            "quote": quote,
            "responses": by_item,
            "now": now_local(),
            "error": None,
            "saved": request.query_params.get("salvo") == "1",
        },
    )


@app.post("/fornecedor/cotacoes/{quote_id}")
async def supplier_save_quote(quote_id: int, request: Request, db: Session = Depends(get_db)):
    user = require_role(request, db, "supplier")
    if user.must_change_password:
        return RedirectResponse("/fornecedor/alterar-senha", 303)
    allowed = db.scalar(select(QuoteSupplier).where(and_(QuoteSupplier.quote_id == quote_id, QuoteSupplier.supplier_id == user.id)))
    quote = db.scalar(select(Quote).where(Quote.id == quote_id).options(selectinload(Quote.items)))
    if not allowed or not quote:
        raise HTTPException(403)
    if now_local() > quote.deadline:
        raise HTTPException(403, "O prazo desta cotação foi encerrado.")
    form = await request.form()
    try:
        item_ids = [item.id for item in quote.items]
        has_response = bool(
            db.scalar(
                select(QuoteResponse.id).where(
                    QuoteResponse.supplier_id == user.id,
                    QuoteResponse.quote_item_id.in_(item_ids),
                ).limit(1)
            )
        ) if item_ids else False
        for item in quote.items:
            price_raw = str(form.get(f"price_{item.id}") or "").strip()
            if not price_raw:
                continue
            has_response = True
            price = parse_decimal(price_raw)
            factor = parse_decimal(form.get(f"factor_{item.id}") or "1")
            if price < 0 or factor <= 0:
                raise ValueError(f"Valor inválido no item {item.description}")
            response = db.scalar(select(QuoteResponse).where(and_(QuoteResponse.quote_item_id == item.id, QuoteResponse.supplier_id == user.id)))
            if not response:
                response = QuoteResponse(quote_item_id=item.id, supplier_id=user.id, offered_price=price)
                db.add(response)
            response.offered_unit = str(form.get(f"unit_{item.id}") or item.base_unit).strip().upper()
            response.base_units_per_offered_unit = factor
            response.offered_price = price
            response.brand = str(form.get(f"brand_{item.id}") or "").strip() or None
            dd = str(form.get(f"delivery_{item.id}") or "").strip()
            response.delivery_days = int(dd) if dd.isdigit() else None
            response.notes = str(form.get(f"notes_{item.id}") or "").strip() or None
        if not has_response:
            raise ValueError("Preencha ao menos um preço antes de enviar a cotação.")
        # O clique em Salvar minha cotação é considerado o envio da proposta.
        allowed.submitted_at = now_local()
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(400, str(exc))
    return RedirectResponse(f"/fornecedor/cotacoes/{quote_id}?salvo=1", 303)
